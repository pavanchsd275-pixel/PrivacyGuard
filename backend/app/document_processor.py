"""
document_processor.py

Step 2 of the PrivacyGuard architecture: Document Processing.

Responsibility: take a file (PDF, DOCX, TXT, or CSV) and return PLAIN TEXT,
preserving structure where reasonably possible. This module does NOT do any
PII detection — it only extracts text so the detection stage has something
to work on.

Design decision: every extractor returns the SAME shape of result, so the
rest of the pipeline (PII detection, anonymization) never needs to know
which file type it came from. This is a common pattern: normalize inputs
early, so downstream logic stays simple.
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Literal

import pymupdf as fitz  # PyMuPDF (new import name; `fitz` alias kept for familiarity)
import docx  # python-docx
import pandas as pd


SupportedType = Literal["pdf", "docx", "txt", "csv"]


@dataclass
class ExtractedDocument:
    """Normalized output of any extractor."""
    text: str                      # full extracted text, structure-preserved as much as possible
    file_type: SupportedType
    source_path: str
    page_count: int | None = None  # only meaningful for PDFs
    metadata: dict = field(default_factory=dict)


class UnsupportedFileTypeError(Exception):
    pass


def _extract_pdf(path: Path) -> ExtractedDocument:
    doc = fitz.open(path)
    pages_text = [page.get_text() for page in doc]
    full_text = "\n\n".join(pages_text)
    return ExtractedDocument(
        text=full_text,
        file_type="pdf",
        source_path=str(path),
        page_count=len(doc),
        metadata={"pages": len(pages_text)},
    )


def _extract_docx(path: Path) -> ExtractedDocument:
    document = docx.Document(path)

    parts = []
    # Paragraphs
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)

    # Tables (resumes/contracts often put contact info in tables!)
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    full_text = "\n".join(parts)
    return ExtractedDocument(
        text=full_text,
        file_type="docx",
        source_path=str(path),
        metadata={"paragraph_count": len(document.paragraphs), "table_count": len(document.tables)},
    )


def _extract_txt(path: Path) -> ExtractedDocument:
    text = path.read_text(encoding="utf-8", errors="replace")
    return ExtractedDocument(
        text=text,
        file_type="txt",
        source_path=str(path),
    )


def _extract_csv(path: Path) -> ExtractedDocument:
    df = pd.read_csv(path, dtype=str, keep_default_na=False)
    # Flatten to text: one "sentence" per row, so PII detection (which works on
    # natural language) has something coherent to scan. We keep column names
    # so context (e.g. "Email: x@y.com") is preserved, which helps NER/regex.
    lines = []
    for _, row in df.iterrows():
        line = ", ".join(f"{col}: {val}" for col, val in row.items() if val)
        lines.append(line)
    full_text = "\n".join(lines)
    return ExtractedDocument(
        text=full_text,
        file_type="csv",
        source_path=str(path),
        metadata={"row_count": len(df), "columns": list(df.columns)},
    )


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".txt": _extract_txt,
    ".csv": _extract_csv,
}


def extract_text(file_path: str | Path) -> ExtractedDocument:
    """
    Main entry point. Give it any supported file path, get back normalized text.

    Raises UnsupportedFileTypeError if the extension isn't handled.
    """
    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix not in _EXTRACTORS:
        raise UnsupportedFileTypeError(
            f"'{suffix}' is not supported. Supported types: {list(_EXTRACTORS.keys())}"
        )

    if not path.exists():
        raise FileNotFoundError(f"No file at {path}")

    return _EXTRACTORS[suffix](path)
